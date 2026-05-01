"""Importa resultados oficiais da Mega-Sena a partir do HTML/HTM da Caixa.

O ZIP histórico publicado pela Caixa contém um arquivo HTML com uma `<table>`
listando os concursos. Cada linha tem (ao menos):
- número do concurso
- data do sorteio (dd/mm/aaaa)
- 6 dezenas

O parser é tolerante: extrai todos os <tr>, lê células <td>, e tenta identificar
as colunas por padrão. Linhas que não casam (cabeçalho, rodapé, prêmios, etc.)
são ignoradas.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path

from django.db import transaction

from resultados.models import ResultadoMegaSena

logger = logging.getLogger(__name__)

_DATA_RE = re.compile(r"\b(\d{2})/(\d{2})/(\d{4})\b")
_DEZENA_RE = re.compile(r"^\s*\d{1,2}\s*$")


class _TableExtractor(HTMLParser):
    """Coleta todas as linhas <tr> como listas de strings de células."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.linhas: list[list[str]] = []
        self._linha: list[str] | None = None
        self._celula: list[str] | None = None

    def handle_starttag(self, tag: str, attrs):
        if tag == "tr":
            self._linha = []
        elif tag in ("td", "th") and self._linha is not None:
            self._celula = []

    def handle_endtag(self, tag: str):
        if tag == "tr" and self._linha is not None:
            if self._linha:
                self.linhas.append(self._linha)
            self._linha = None
        elif tag in ("td", "th") and self._celula is not None:
            self._linha.append("".join(self._celula).strip())
            self._celula = None

    def handle_data(self, data: str):
        if self._celula is not None:
            self._celula.append(data)


@dataclass(frozen=True)
class ResultadoExtraido:
    concurso: int
    data: date
    dezenas: tuple[int, int, int, int, int, int]


def parse_html_megasena(html: str) -> list[ResultadoExtraido]:
    parser = _TableExtractor()
    parser.feed(html)

    resultados: list[ResultadoExtraido] = []
    for linha in parser.linhas:
        # localiza a primeira célula numérica como concurso e a primeira data dd/mm/aaaa
        concurso: int | None = None
        data: date | None = None
        dezenas: list[int] = []

        for celula in linha:
            txt = celula.strip()
            if not txt:
                continue
            if concurso is None and txt.isdigit() and 1 <= len(txt) <= 5:
                concurso = int(txt)
                continue
            if data is None:
                m = _DATA_RE.search(txt)
                if m:
                    try:
                        data = datetime.strptime(
                            f"{m.group(1)}/{m.group(2)}/{m.group(3)}", "%d/%m/%Y"
                        ).date()
                    except ValueError:
                        pass
                    continue
            if _DEZENA_RE.match(txt):
                n = int(txt)
                if 1 <= n <= 60 and len(dezenas) < 6:
                    dezenas.append(n)

        if concurso is not None and data is not None and len(dezenas) == 6:
            resultados.append(
                ResultadoExtraido(concurso=concurso, data=data, dezenas=tuple(dezenas))
            )

    return resultados


def importar_de_arquivo_html(path: Path) -> dict:
    """Lê um HTML/HTM e cria/atualiza ResultadoMegaSena. Idempotente por concurso."""
    html = path.read_text(encoding="latin-1", errors="ignore")
    extraidos = parse_html_megasena(html)
    return _persistir(extraidos)


def parse_xlsx_megasena(path: Path) -> list[ResultadoExtraido]:
    """Lê o XLSX oficial da Caixa (aba "MEGA SENA") e devolve os resultados.

    Formato esperado (header na primeira linha):
        Concurso | Data do Sorteio | Bola1 | Bola2 | Bola3 | Bola4 | Bola5 | Bola6 | ...
    """
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), data_only=True)
    # Usa a primeira aba; aceita variações de capitalização do nome.
    ws = wb.worksheets[0]

    resultados: list[ResultadoExtraido] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # cabeçalho
        if not row or row[0] in (None, ""):
            continue
        try:
            concurso = int(row[0])
        except (TypeError, ValueError):
            continue

        data_raw = row[1]
        if isinstance(data_raw, date) and not isinstance(data_raw, datetime):
            data = data_raw
        elif isinstance(data_raw, datetime):
            data = data_raw.date()
        elif isinstance(data_raw, str):
            m = _DATA_RE.search(data_raw)
            if not m:
                continue
            data = _try_data(m.group(1), m.group(2), m.group(3))
            if data is None:
                continue
        else:
            continue

        try:
            dezenas = tuple(int(row[c]) for c in range(2, 8))
        except (TypeError, ValueError):
            continue
        if any(not (1 <= d <= 60) for d in dezenas):
            continue

        resultados.append(ResultadoExtraido(concurso=concurso, data=data, dezenas=dezenas))

    return resultados


def importar_de_arquivo_xlsx(path: Path) -> dict:
    extraidos = parse_xlsx_megasena(path)
    return _persistir(extraidos)


def _try_data(d: str, m: str, y: str) -> date | None:
    try:
        return datetime.strptime(f"{d}/{m}/{y}", "%d/%m/%Y").date()
    except ValueError:
        return None


def _persistir(extraidos: list[ResultadoExtraido]) -> dict:
    criados = atualizados = 0
    with transaction.atomic():
        for r in extraidos:
            obj, created = ResultadoMegaSena.objects.update_or_create(
                concurso=r.concurso,
                defaults={
                    "data_apuracao": r.data,
                    "dezena_1": r.dezenas[0],
                    "dezena_2": r.dezenas[1],
                    "dezena_3": r.dezenas[2],
                    "dezena_4": r.dezenas[3],
                    "dezena_5": r.dezenas[4],
                    "dezena_6": r.dezenas[5],
                },
            )
            if created:
                criados += 1
            else:
                atualizados += 1
    return {
        "encontrados": len(extraidos),
        "criados": criados,
        "atualizados": atualizados,
    }
